"""Build excel/assumptions_and_sensitivity.xlsx: assumption cells, data-derived inputs, live formulas.

Colour code: yellow cells with blue text are ASSUMPTION inputs; grey cells are DATA-DERIVED numbers pasted from the
pipeline; everything else is a formula. Change an assumption and the finding, tornado and two-way table recompute.
Legacy function names (NORMSDIST, NORMDIST, NORMSINV, TINV) are used on purpose: they need no _xlfn prefix, so the
file calculates identically in every Excel version and in the COM check.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

from .cost_model import BASE, RANGES

ASSUME_FILL = PatternFill("solid", fgColor="FFF2CC")
DATA_FILL = PatternFill("solid", fgColor="E7E6E6")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
BLUE = Font(color="0000FF")
HEAD_FONT = Font(bold=True, color="FFFFFF")
FIRST_ROW = 4

NAMES = {  # defined name -> Assumptions cell
    "LeadTime_L": "B4", "ReviewPeriod_R": "B5", "HoldRate_h": "B6", "Margin_m": "B7", "OrderCost_K": "B8",
    "UnitPrice_p": "B9", "ServiceLevel": "B10", "ServiceZ": "B11", "UnitCost_c": "B12", "LossFn_Lz": "B13",
    "K_factor": "B14",
}
SUMMARY = {"mean": "K4", "sd": "K5", "n": "K6", "t_lo": "K7", "t_hi": "K8", "boot_lo": "K9", "boot_hi": "K10",
           "mean_snaive": "K11", "mean_delta": "K12"}
CHAIN_FINDING = "K27"
CHECK_CELL = "B10"


def k_formula(L="LeadTime_L", R="ReviewPeriod_R", h="HoldRate_h", m="Margin_m", p="UnitPrice_p", z="ServiceZ") -> str:
    """Excel text of k = sqrt(R+L) * (z p (1-m) h + (365/R) Lz(z) m p); arguments are cell refs or names."""
    lz = f"(NORMDIST({z},0,1,FALSE)-{z}*(1-NORMSDIST({z})))"
    return f"SQRT({R}+{L})*({z}*{p}*(1-{m})*{h}+(365/{R})*{lz}*{m}*{p})"


def _header(ws, row: int, labels: list[str]) -> None:
    for j, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=j, value=label)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(wrap_text=True, vertical="center")


def _widths(ws, widths: dict[str, float]) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _readme(wb: Workbook, is_demo: bool) -> None:
    ws = wb.active
    ws.title = "README"
    lines = [
        ("Cost of forecast error: assumptions and sensitivity", True),
        ("SYNTHETIC DEMO DATA: the data-derived inputs come from make_demo_data.py, not from Rossmann. "
         "This file only tests the pipeline." if is_demo else
         "Data-derived inputs come from the Rossmann Store Sales pipeline (Kaggle, not redistributed).", True),
        ("", False),
        ("What is real and what is assumed", True),
        ("Rossmann has NO stock, orders, lead times, costs or unit prices. Every yellow cell on the Assumptions sheet is "
         "an ASSUMPTION with a low/high range. The grey cells on Data_inputs are forecast-error sizes measured on the "
         "validation months (2015-01-01 to 2015-06-24).", False),
        ("Sales is daily turnover, not units. One 'unit' here is EUR 1 of turnover (unit price p = 1 is a normalisation, "
         "not an estimate). The currency itself is the community reading, not an official statement.", False),
        ("", False),
        ("The one quantified finding", True),
        ("Extra expected cost per store-year when the daily forecast error behaves like a promotion day instead of an "
         "ordinary day, under an order-up-to policy with safety stock sized from the forecast error. Descriptive and "
         "assumption-driven: it is not a measured saving.", False),
        ("Sheets: Assumptions (inputs) > Data_inputs (measured errors) > Finding (formulas, CI) > Tornado > TwoWay > Checks.", False),
        ("Errors are assumed uncorrelated across days when scaled by sqrt(R+L); the calibration table reports "
         "Ljung-Box so the optimism can be judged.", False),
    ]
    for i, (text, bold) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font = Font(bold=bold, size=12 if i == 1 else 11)
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 110


def _assumptions(wb: Workbook) -> None:
    ws = wb.create_sheet("Assumptions")
    ws["A1"] = "Assumptions (yellow = ASSUMPTION, change me)"
    ws["A1"].font = Font(bold=True, size=12)
    _header(ws, 3, ["Parameter", "Value", "Unit", "Low", "High", "Label", "Why this value"])
    rows = [
        ("Lead time L", BASE.lead_time_days, "days", *RANGES["lead_time_days"], "ASSUMPTION",
         "Typical depot-to-store lead time; Rossmann provides none"),
        ("Review period R", BASE.review_period_days, "days", *RANGES["review_period_days"], "ASSUMPTION",
         "Weekly store ordering cycle"),
        ("Holding cost rate h", BASE.holding_rate, "per year", *RANGES["holding_rate"], "ASSUMPTION",
         "Textbook retail carrying-cost range 20 to 30 percent"),
        ("Gross margin m (lost-sale cost)", BASE.margin, "share of price", *RANGES["margin"], "ASSUMPTION",
         "Foregone margin proxy for one lost sale; no margin data exists"),
        ("Order cost K", BASE.order_cost, "EUR per order", *RANGES["order_cost"], "ASSUMPTION",
         "Not part of the finding; shown for completeness"),
        ("Unit price p", BASE.unit_price, "EUR per unit", None, None, "ASSUMPTION",
         "Normalisation: 1 unit = EUR 1 of turnover"),
        ("Cycle service level", BASE.service_level, "share", None, None, "ASSUMPTION", "Policy choice, not a cost"),
    ]
    for i, row in enumerate(rows, 4):
        for j, v in enumerate(row, 1):
            ws.cell(row=i, column=j, value=v)
        for col in (2,):
            ws.cell(row=i, column=col).fill, ws.cell(row=i, column=col).font = ASSUME_FILL, BLUE
    derived = [
        (11, "z = NORMSINV(service level)", "=NORMSINV(ServiceLevel)"),
        (12, "Unit cost c = p (1 - m)", "=UnitPrice_p*(1-Margin_m)"),
        (13, "Loss function Lz = phi(z) - z (1 - Phi(z))", "=NORMDIST(ServiceZ,0,1,FALSE)-ServiceZ*(1-NORMSDIST(ServiceZ))"),
        (14, "k = extra EUR per year per EUR 1 of daily RMSE", "=" + k_formula()),
    ]
    for r, label, formula in derived:
        ws.cell(row=r, column=1, value=label)
        ws.cell(row=r, column=2, value=formula)
        ws.cell(row=r, column=6, value="DERIVED (formula)")
    for name, ref in NAMES.items():
        wb.defined_names[name] = DefinedName(name, attr_text=f"Assumptions!${ref[0]}${ref[1:]}")
    _widths(ws, {"A": 46, "B": 12, "C": 16, "D": 9, "E": 9, "F": 18, "G": 60})


def _data_inputs(wb: Workbook, ml: pd.DataFrame, sn: pd.DataFrame) -> int:
    ws = wb.create_sheet("Data_inputs")
    ws["A1"] = "DATA-DERIVED: RMSE of the daily forecast error per store, validation months, horizons 1-14"
    ws["A1"].font = Font(bold=True, size=12)
    _header(ws, 3, ["store", "ML rmse promo", "ML rmse non-promo", "n promo days", "n non-promo days",
                    "mean daily sales", "seasonal-naive rmse promo", "seasonal-naive rmse non-promo"])
    merged = ml.merge(sn[["store", "rmse_promo", "rmse_nonpromo"]], on="store", suffixes=("", "_sn"))
    for i, r in enumerate(merged.itertuples(index=False), FIRST_ROW):
        vals = [int(r.store), r.rmse_promo, r.rmse_nonpromo, int(r.n_promo), int(r.n_nonpromo), r.mean_daily_sales,
                r.rmse_promo_sn, r.rmse_nonpromo_sn]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=i, column=j, value=float(v) if j > 1 else v)
            c.fill = DATA_FILL
    ws.freeze_panes = "B4"
    _widths(ws, {get_column_letter(i): 16 for i in range(1, 9)})
    return FIRST_ROW + len(merged) - 1


def _finding(wb: Workbook, last: int, boot: tuple[float, float, float]) -> None:
    ws = wb.create_sheet("Finding")
    ws["A1"] = "Finding: extra cost per store-year, promo-like error minus ordinary-day error (EUR)"
    ws["A1"].font = Font(bold=True, size=12)
    _header(ws, 3, ["store", "delta rmse (ML)", "Extra cost promo (ML)", "Extra cost non-promo (ML)",
                    "Finding (ML)", "delta rmse (seasonal naive)", "Finding (seasonal naive)"])
    for i in range(FIRST_ROW, last + 1):
        ws[f"A{i}"] = f"=Data_inputs!A{i}"
        ws[f"B{i}"] = f"=Data_inputs!B{i}-Data_inputs!C{i}"
        ws[f"C{i}"] = f"=Data_inputs!B{i}*K_factor"
        ws[f"D{i}"] = f"=Data_inputs!C{i}*K_factor"
        ws[f"E{i}"] = f"=C{i}-D{i}"
        ws[f"F{i}"] = f"=Data_inputs!G{i}-Data_inputs!H{i}"
        ws[f"G{i}"] = f"=F{i}*K_factor"
    rng = f"E{FIRST_ROW}:E{last}"
    summary = [
        ("Mean finding, ML forecast (EUR per store-year)", f"=AVERAGE({rng})"),
        ("SD across stores", f"=STDEV({rng})"),
        ("Stores (n)", f"=COUNT({rng})"),
        ("95% CI low (t, across stores)", "=K4-TINV(0.05,K6-1)*K5/SQRT(K6)"),
        ("95% CI high (t, across stores)", "=K4+TINV(0.05,K6-1)*K5/SQRT(K6)"),
        ("95% CI low (Python store bootstrap, static)", boot[1]),
        ("95% CI high (Python store bootstrap, static)", boot[2]),
        ("Mean finding, seasonal-naive forecast", f"=AVERAGE(G{FIRST_ROW}:G{last})"),
        ("Mean delta rmse (ML), EUR per day", f"=AVERAGE(B{FIRST_ROW}:B{last})"),
    ]
    for r, (label, value) in enumerate(summary, 4):
        ws[f"J{r}"], ws[f"K{r}"] = label, value
    ws["K9"].fill = ws["K10"].fill = DATA_FILL
    ws["J14"] = "Literal formula chain on the average store (ML), one line per step"
    ws["J14"].font = Font(bold=True)
    chain = [
        (15, "sigma promo = mean rmse promo", f"=AVERAGE(Data_inputs!B{FIRST_ROW}:B{last})"),
        (16, "sigma non-promo = mean rmse non-promo", f"=AVERAGE(Data_inputs!C{FIRST_ROW}:C{last})"),
        (17, "SS promo = z sigma sqrt(R+L)", "=ServiceZ*K15*SQRT(ReviewPeriod_R+LeadTime_L)"),
        (18, "Holding promo = SS c h", "=K17*UnitCost_c*HoldRate_h"),
        (19, "Shortage per cycle promo = sigma sqrt(R+L) Lz", "=K15*SQRT(ReviewPeriod_R+LeadTime_L)*LossFn_Lz"),
        (20, "Shortage cost promo = (365/R) * that * m p", "=(365/ReviewPeriod_R)*K19*Margin_m*UnitPrice_p"),
        (21, "Extra promo = Holding + Shortage", "=K18+K20"),
        (22, "SS non-promo", "=ServiceZ*K16*SQRT(ReviewPeriod_R+LeadTime_L)"),
        (23, "Holding non-promo", "=K22*UnitCost_c*HoldRate_h"),
        (24, "Shortage per cycle non-promo", "=K16*SQRT(ReviewPeriod_R+LeadTime_L)*LossFn_Lz"),
        (25, "Shortage cost non-promo", "=(365/ReviewPeriod_R)*K24*Margin_m*UnitPrice_p"),
        (26, "Extra non-promo", "=K23+K25"),
        (27, "FINDING = Extra promo - Extra non-promo (EUR per store-year)", "=K21-K26"),
    ]
    for r, label, formula in chain:
        ws[f"J{r}"], ws[f"K{r}"] = label, formula
    ws["J27"].font = ws["K27"].font = Font(bold=True)
    ws.freeze_panes = "B4"
    _widths(ws, {"A": 9, "B": 16, "C": 18, "D": 18, "E": 16, "F": 20, "G": 20, "J": 58, "K": 16})


def _tornado(wb: Workbook) -> None:
    ws = wb.create_sheet("Tornado")
    ws["A1"] = "One-way sensitivity of the finding (EUR per store-year); order cost K does not enter the finding by construction"
    ws["A1"].font = Font(bold=True, size=12)
    _header(ws, 3, ["Parameter", "Low", "High", "Finding at low", "Finding at high", "Swing", "Rank"])
    base = {"L": "LeadTime_L", "R": "ReviewPeriod_R", "h": "HoldRate_h", "m": "Margin_m", "p": "UnitPrice_p", "z": "ServiceZ"}
    params = [("Lead time L (days)", "L", "lead_time_days"), ("Review period R (days)", "R", "review_period_days"),
              ("Holding rate h", "h", "holding_rate"), ("Margin m", "m", "margin"), ("Order cost K (EUR)", None, "order_cost")]
    for i, (label, key, rng_key) in enumerate(params, 4):
        lo, hi = RANGES[rng_key]
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = label, lo, hi
        for col, src in (("D", f"B{i}"), ("E", f"C{i}")):
            args = dict(base)
            if key:
                args[key] = src
            ws[f"{col}{i}"] = f"=Finding!$K$12*{k_formula(args['L'], args['R'], args['h'], args['m'], args['p'], args['z'])}"
        ws[f"F{i}"] = f"=ABS(E{i}-D{i})"
        ws[f"G{i}"] = f"=RANK(F{i},$F$4:$F$8)"
        ws[f"B{i}"].fill = ws[f"C{i}"].fill = ASSUME_FILL
    _widths(ws, {"A": 26, "B": 10, "C": 10, "D": 16, "E": 16, "F": 12, "G": 8})


def _two_way(wb: Workbook) -> None:
    ws = wb.create_sheet("TwoWay")
    ws["A1"] = "Finding (EUR per store-year): lead time L (rows) x margin m (columns)"
    ws["A1"].font = Font(bold=True, size=12)
    ls, ms = [2, 3, 5, 7, 10], [0.15, 0.225, 0.30, 0.40, 0.50]
    ws["A3"] = "L \\ m"
    for j, m in enumerate(ms, 2):
        c = ws.cell(row=3, column=j, value=m)
        c.fill, c.font = ASSUME_FILL, BLUE
    for i, L in enumerate(ls, 4):
        c = ws.cell(row=i, column=1, value=L)
        c.fill, c.font = ASSUME_FILL, BLUE
        for j in range(2, 2 + len(ms)):
            col = get_column_letter(j)
            ws.cell(row=i, column=j, value=f"=Finding!$K$12*{k_formula(L=f'$A{i}', m=f'{col}$3')}")
    _widths(ws, {"A": 10, **{get_column_letter(j): 12 for j in range(2, 8)}})


def _checks(wb: Workbook, last: int) -> None:
    ws = wb.create_sheet("Checks")
    ws["A1"] = "Sanity limit cases (live)"
    ws["A1"].font = Font(bold=True, size=12)
    _header(ws, 3, ["Check", "Value", "Pass"])
    checks = [
        ("Zero error gives zero extra cost", "=0*K_factor", "=ABS(B4)<1E-12"),
        ("Equal promo and ordinary error gives zero finding", "=(100-100)*K_factor", "=ABS(B5)<1E-12"),
        ("Extra cost is increasing in sigma (2s minus s)", "=200*K_factor-100*K_factor", "=B6>0"),
        ("Formula chain equals mean of per-store findings", f"=Finding!{CHAIN_FINDING}-Finding!{SUMMARY['mean']}",
         "=ABS(B7)<1E-6*MAX(1,ABS(Finding!K4))"),
        ("No error cells in the finding column", f"=SUMPRODUCT(--ISERROR(Finding!E{FIRST_ROW}:E{last}))", "=B8=0"),
    ]
    for i, (label, value, passed) in enumerate(checks, 4):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"] = label, value, passed
    ws["A10"] = "Overall"
    ws[CHECK_CELL] = '=IF(AND(C4:C8),"ALL CHECKS PASS","CHECK FAILED")'
    _widths(ws, {"A": 54, "B": 18, "C": 10})


def build_workbook(path: Path, ml_inputs: pd.DataFrame, sn_inputs: pd.DataFrame, boot: tuple[float, float, float],
                   is_demo: bool) -> Path:
    wb = Workbook()
    _readme(wb, is_demo)
    _assumptions(wb)
    last = _data_inputs(wb, ml_inputs, sn_inputs)
    _finding(wb, last, boot)
    _tornado(wb)
    _two_way(wb)
    _checks(wb, last)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
